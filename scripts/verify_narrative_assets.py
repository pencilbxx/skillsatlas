"""Verify the 10 Sep narrative alignment and generated artefacts."""

from pathlib import Path

import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"


def require(text: str, needle: str, label: str) -> None:
    if needle not in text:
        raise AssertionError(f"{label}: missing {needle!r}")


def forbid(text: str, needle: str, label: str) -> None:
    if needle.lower() in text.lower():
        raise AssertionError(f"{label}: stale or forbidden {needle!r}")


def docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def pdf_text(path: Path) -> str:
    doc = pymupdf.open(path)
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def main() -> None:
    canonical = (DOCS / "pitch-slide-text.md").read_text(encoding="utf-8")
    for marker in (
        "A list gives you options. SkillsAtlas gives you proof.",
        "SkillsAtlas Interview Board",
        "Don",
        "Sophia",
        "Bridget",
        "Andrew + Sri Karan",
        "clickable prototype using a prepared practice CV",
        "WhatsApp is removed from the active pitch and demo",
    ):
        require(canonical, marker, "canonical Team Draft")

    word = docx_text(DOCS / "SkillsAtlas_Hub_Pitch_Slide_Text_TEAM_DRAFT.docx")
    for marker in (
        "A list gives you options. SkillsAtlas gives you proof.",
        "Don",
        "Sophia",
        "Bridget",
        "Andrew + Sri Karan",
        "Files that control the active experience",
        "Commenting reminder",
    ):
        require(word, marker, "generated Word Team Draft")

    canvas = pdf_text(DOCS / "canvas" / "lean-canvas-working.pdf")
    for marker in (
        "A list gives you options. SkillsAtlas gives you proof.",
        "THE WORKER DOES NOT PAY",
        "recruiter-alone CRM search",
        "WORKING · OPEN ITEMS MARKED",
    ):
        require(canvas, marker, "Lean Canvas rev 7")
    for stale in ("Everyone else hands you a list", "silent weeks", "REVISION 6"):
        forbid(canvas, stale, "Lean Canvas rev 7")

    html = (DOCS / "wireframes" / "demo.html").read_text(encoding="utf-8")
    js = (DOCS / "wireframes" / "demo.js").read_text(encoding="utf-8")
    active_demo = html + "\n" + js
    for marker in (
        "A list gives you options.",
        "SkillsAtlas Interview Board",
        "Confirm stronger wording",
        "Illustrative learning option",
        'evidenceState = "proposed"',
        'evidenceState = "enriched"',
        "confirmBtn.disabled = true",
        'if (evidenceState !== "unconfirmed") return',
        "seeded preview",
        "original CV wording confirmed",
    ):
        require(active_demo, marker, "active demo")
    for stale in (
        "WhatsApp",
        "cheat-board",
        "Predicted questions",
        "See you in three weeks",
        "AI practice helper",
        "Scripted practice helper",
    ):
        forbid(active_demo, stale, "active demo")
    forbid(js, "if (index === 1) resetEvidence()", "active demo navigation")

    workplan = (DOCS / "team-workplan-10sep.md").read_text(encoding="utf-8")
    for marker in ("Don", "Sophia", "Bridget", "Andrew + Sri Karan", "feature/interview-practice-helper"):
        require(workplan, marker, "team workplan")

    print("Narrative, Word Team Draft, Lean Canvas, demo, and workplan checks passed.")


if __name__ == "__main__":
    main()
