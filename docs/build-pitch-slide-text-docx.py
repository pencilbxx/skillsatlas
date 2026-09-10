"""Build the Hub pitch Word Team Draft from the canonical Markdown.

Source: docs/pitch-slide-text.md
Output: docs/SkillsAtlas_Hub_Pitch_Slide_Text_TEAM_DRAFT.docx

The Markdown file is the source of truth. Do not duplicate pitch copy in this script.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "pitch-slide-text.md"
OUT = ROOT / "SkillsAtlas_Hub_Pitch_Slide_Text_TEAM_DRAFT.docx"

NAVY = RGBColor(0x10, 0x2B, 0x3F)
CORAL = RGBColor(0xED, 0x6A, 0x4A)
TEAL = RGBColor(0x1E, 0x5D, 0x5E)
BODY = RGBColor(0x18, 0x31, 0x43)
MUTED = RGBColor(0x49, 0x60, 0x6D)
CREAM = "F7F4EF"
HEADER_FILL = "102B3F"
QUOTE_FILL = "EEF5F1"
ALT_FILL = "F7F4EF"


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, *, size=10.5, bold=False, italic=False, color=BODY, font="Aptos") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_inline(paragraph, text: str, *, size=10.5, color=BODY, italic_all=False) -> None:
    """Render a small, deliberate subset of Markdown inline formatting."""
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            style_run(run, size=size, color=color, italic=italic_all)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            style_run(run, size=size, color=color, bold=True, italic=italic_all)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            style_run(run, size=size - 0.5, color=TEAL, font="Aptos Mono")
        else:
            run = paragraph.add_run(token[1:-1])
            style_run(run, size=size, color=color, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        style_run(run, size=size, color=color, italic=italic_all)


def add_paragraph(container, text: str, *, size=10.5, color=BODY, bold=False, italic=False, before=0, after=5):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if bold:
        run = p.add_run(text)
        style_run(run, size=size, color=color, bold=True, italic=italic)
    else:
        add_inline(p, text, size=size, color=color, italic_all=italic)
    return p


def add_heading(doc: Document, text: str, level: int) -> None:
    if level == 1:
        size, color, before, after = 24, NAVY, 0, 8
    elif level == 2:
        size, color, before, after = 16, NAVY, 12, 5
    else:
        size, color, before, after = 12, CORAL, 8, 3
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    style_run(run, size=size, bold=True, color=color)


def add_quote(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, QUOTE_FILL)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, size=11.5, color=NAVY, italic_all=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell)
            if r_idx == 0:
                shade(cell, HEADER_FILL)
                color = RGBColor(0xFF, 0xFF, 0xFF)
                bold = True
            else:
                if r_idx % 2 == 0:
                    shade(cell, ALT_FILL)
                color = BODY
                bold = False
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text = row[c_idx] if c_idx < len(row) else ""
            if bold:
                run = p.add_run(text.replace("**", ""))
                style_run(run, size=9.2, bold=True, color=color)
            else:
                add_inline(p, text, size=9.2, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("SkillsAtlas · Hub pitch team draft · canonical copy")
    style_run(run, size=8, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Hub 14 Sep 2026 · slides freeze 13 Sep, 2pm")
    style_run(run, size=8, color=MUTED)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BODY
    return doc


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = setup_document()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 3)
        elif stripped.startswith("## "):
            if stripped[3:] == "Files that control the active experience":
                doc.add_page_break()
            add_heading(doc, stripped[3:], 2)
        elif stripped.startswith("# "):
            add_heading(doc, stripped[2:], 1)
        elif stripped.startswith("> "):
            quote_lines = [stripped[2:]]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("> "):
                quote_lines.append(lines[i].strip()[2:])
                i += 1
            add_quote(doc, " ".join(quote_lines))
            continue
        elif re.match(r"^- \[[ xX]\] ", stripped):
            checked = stripped[3].lower() == "x"
            text = stripped[6:]
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, ("☒ " if checked else "☐ ") + text, size=10.2)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, stripped[2:], size=10.2)
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\d+\. ", "", stripped), size=10.2)
        elif stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p_pr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D9CFC3")
            borders.append(bottom)
            p_pr.append(borders)
        else:
            paragraph_lines = [stripped]
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if not nxt:
                    break
                if nxt.startswith(("#", "- ", "> ", "|")) or re.match(r"^\d+\. ", nxt):
                    break
                paragraph_lines.append(nxt)
                i += 1
            add_paragraph(doc, " ".join(paragraph_lines))
            continue
        i += 1

    add_heading(doc, "Commenting reminder", 2)
    add_paragraph(
        doc,
        "Turn on Review → Track Changes in Word. Resolve the narrative before slide design. "
        "If a line changes here, make the same change in docs/pitch-slide-text.md and regenerate this file.",
        color=MUTED,
        italic=True,
    )
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print("wrote", build())
